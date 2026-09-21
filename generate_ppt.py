"""Generate a professional widescreen PPTX presentation with embedded output images.
Title: AI-Powered Market Trend and Consumer Sentiment Forecaster
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path="AI_Powered_Market_Trend_and_Consumer_Sentiment_Forecaster.pptx"):
    prs = Presentation()
    # Widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette: Modern Navy / Indigo Tech
    BG_COLOR = RGBColor(15, 23, 42)        # Slate 900
    CARD_BG = RGBColor(30, 41, 59)         # Slate 800
    ACCENT_INDIGO = RGBColor(99, 102, 241) # Indigo 500
    ACCENT_CYAN = RGBColor(6, 182, 212)    # Cyan 500
    TEXT_WHITE = RGBColor(248, 250, 252)   # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)   # Slate 400
    TEXT_BODY = RGBColor(226, 232, 240)    # Slate 200
    ACCENT_EMERALD = RGBColor(16, 185, 129)# Emerald 500
    ACCENT_ROSE = RGBColor(244, 63, 94)    # Rose 500

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, category_tag, title_text):
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.40), Inches(11.5), Inches(0.35))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_CYAN

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.70), Inches(11.5), Inches(0.75))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(25)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------
    # SLIDE 1: Title Page
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide1)

    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.4), Inches(0.15), Inches(4.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_INDIGO
    bar.line.fill.background()

    t_box = slide1.shapes.add_textbox(Inches(1.4), Inches(1.3), Inches(10.8), Inches(4.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "FINAL YEAR MAJOR PROJECT PRESENTATION"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN
    p0.space_after = Pt(16)

    p1 = tf1.add_paragraph()
    p1.text = "AI-Powered Market Trend and Consumer Sentiment Forecaster"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "A Deep Learning & Unsupervised NLP Intelligence Framework with Aspect-Based Sentiment Analysis, BERTopic Semantic Modeling, and SentiTSMixer Sales Forecasting"
    p2.font.size = Pt(16)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(36)

    p3 = tf1.add_paragraph()
    p3.text = "Domain: Artificial Intelligence | Deep Time-Series | Natural Language Processing"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = ACCENT_EMERALD

    # -------------------------------------------------------------
    # SLIDE 2: Abstract
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide2)
    add_header(slide2, "Executive Overview", "Abstract")

    card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.1))
    card2.fill.solid()
    card2.fill.fore_color.rgb = CARD_BG
    card2.line.color.rgb = ACCENT_INDIGO
    card2.line.width = Pt(1)

    tf2 = card2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.5)
    tf2.margin_right = Inches(0.5)
    tf2.margin_top = Inches(0.4)

    abstract_paras = [
        ("Background: ", "The explosive growth of e-commerce platforms, social media, and digital news generates immense volumes of unstructured customer feedback, ratings, and market signals."),
        ("Core Problem: ", "Traditional sales forecasting relies purely on historical numeric transactions, ignoring real-time sentiment shifts, review helpfulness votes, and hidden discussion themes."),
        ("Proposed Solution: ", "This project presents an end-to-end AI framework amalgamating Customer Ratings (R), Normalized Review Sentiment (S via BERT & VADER), Review Helpfulness (H), and Cumulative Mixture (SC) into the Align Stage of SentiTSMixer."),
        ("NLP & RAG Integration: ", "Employs BERTopic and Latent Dirichlet Allocation (LDA) for unsupervised trend and topic discovery, combined with FAISS dense vector search and LLMs for Retrieval-Augmented Generation (RAG)."),
        ("Results & Impact: ", "Delivers a 65% to 99% reduction in forecasting error (MAPE, MSE) over ARIMA and LSTM baselines, dampening supply chain bullwhip effects and supporting real-time data-driven strategic decisions.")
    ]

    for i, (bold_prefix, body) in enumerate(abstract_paras):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.space_after = Pt(12)
        run1 = p.add_run()
        run1.text = bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(15)
        run1.font.color.rgb = ACCENT_CYAN

        run2 = p.add_run()
        run2.text = body
        run2.font.size = Pt(14)
        run2.font.color.rgb = TEXT_BODY

    # -------------------------------------------------------------
    # SLIDE 3: Introduction
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide3)
    add_header(slide3, "System Foundations", "Introduction")

    intro_points = [
        ("The Surge of Digital Consumer Voice", "E-commerce and social platforms (Amazon, Twitter/X, Reddit) have made customer reviews the primary driver of product demand. Consumer sentiment shifts weeks before appearing in backward-looking sales records."),
        ("Multi-Dimensional Customer Feedback", "Customer feedback encompasses discrete numeric star ratings (1–5), detailed unstructured textual reviews, and community upvotes (helpfulness likes/dislikes) that validate review truthfulness."),
        ("The Sales Forecasting Imperative", "Accurate business demand forecasting is crucial to balance market supply and production. Over-forecasting ties up capital and causes inventory wastage; under-forecasting leads to stockouts, lost revenue, and brand defection."),
        ("Next-Gen AI Forecasting Paradigm", "Bridging the gap between Natural Language Processing (BERT/VADER, BERTopic, RAG) and Deep Multivariate Time-Series (TSMixer architecture) to construct an intelligent real-time forecasting dashboard.")
    ]

    for idx, (title, desc) in enumerate(intro_points):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.7 + row * 2.6)

        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.35))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_INDIGO if idx % 2 == 0 else ACCENT_CYAN
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p1 = tf.paragraphs[0]
        p1.text = f"0{idx+1}.  {title}"
        p1.font.bold = True
        p1.font.size = Pt(15)
        p1.font.color.rgb = ACCENT_CYAN
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_BODY

    # -------------------------------------------------------------
    # SLIDE 4: Literature Survey (Part 1)
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide4)
    add_header(slide4, "Background Studies", "Literature Survey (Part 1: Foundational Time-Series & Sentiment)")

    papers_p1 = [
        ("Paper 1: Classical Statistical & Neural Time-Series Models",
         "Authors: Taylor & Letham (2018), Fattah et al. (2018), Yenidogan et al. (2018)",
         [
             "Investigated ARIMA, SARIMA, and LSTM architectures for demand and stock prediction.",
             "Findings: ARIMA requires strict stationarity and fails on nonlinear trends; SARIMA captures seasonal periods but breaks under dynamic seasonality.",
             "LSTMs mitigate vanishing gradients with memory gates, yet struggle to fuse auxiliary unstructured text features efficiently without high computational overhead."
         ]),
        ("Paper 2: Customer Sentiment & Lexical Rating Integration",
         "Authors: Ghosh et al. (2021, 2024), Hutto & Gilbert (2014 - VADER)",
         [
             "Developed VADER rule-based sentiment reasoning and explored sales forecasting of overrated products.",
             "Demonstrated that customer ratings are discrete (1-5) and lack granularity; textual reviews provide continuous nuance.",
             "Gap Identified: Models omitted review helpfulness community votes, allowing fake reviews and outlier noise to inflate forecast MSE."
         ])
    ]

    for idx, (p_title, p_auth, p_bullets) in enumerate(papers_p1):
        y = Inches(1.7 + idx * 2.65)
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_INDIGO
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.4)
        tf.margin_right = Inches(0.4)
        tf.margin_top = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = p_title
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = ACCENT_CYAN

        p_a = tf.add_paragraph()
        p_a.text = p_auth
        p_a.font.size = Pt(11)
        p_a.font.color.rgb = TEXT_MUTED
        p_a.space_after = Pt(6)

        for bullet in p_bullets:
            pb = tf.add_paragraph()
            pb.text = f"•  {bullet}"
            pb.font.size = Pt(12)
            pb.font.color.rgb = TEXT_BODY
            pb.space_after = Pt(2)

    # -------------------------------------------------------------
    # SLIDE 5: Literature Survey (Part 2)
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide5)
    add_header(slide5, "Background Studies", "Literature Survey (Part 2: Modern Transformers & TSMixer)")

    papers_p2 = [
        ("Paper 3: TSMixer: An All-MLP Architecture for Time Series",
         "Authors: Si-An Chen, Chun-Liang Li, Nate Yoder, Sercan O. Arik, Tomas Pfister (Google Research, 2023)",
         [
             "Introduced TSMixer, replacing complex self-attention with Time-Mixing MLPs and Feature-Mixing MLPs.",
             "Achieved superior accuracy with 2D Layer Normalization and residual connections at a fraction of Transformer compute.",
             "Demonstrated the 'Align Stage' concept for projecting auxiliary static and time-varying auxiliary features into shared time-series representations."
         ]),
        ("Paper 4: SentiTSMixer: Specific Sales Forecasting with Customer Sentiment",
         "Authors: Ghosh, Das, Roy, Bhattacharjee, Cortesi, Sen (IEEE Access, Vol. 13, 2025)",
         [
             "Direct grounding paper for this project. Integrated BERT/VADER sentiment, ratings, and helpfulness upvotes into TSMixer.",
             "Formulated Equation 8 (VADER 0-1 alignment), Eq. 9 (1-5 sentiment scaling), Eq. 10 (helpfulness normalization), and Eq. 11 (mixture SC).",
             "Demonstrated Random Forest MDI feature importance and proved 65% to 99% error reduction across Amazon datasets."
         ])
    ]

    for idx, (p_title, p_auth, p_bullets) in enumerate(papers_p2):
        y = Inches(1.7 + idx * 2.65)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_CYAN
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.4)
        tf.margin_right = Inches(0.4)
        tf.margin_top = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = p_title
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = ACCENT_INDIGO if idx == 0 else ACCENT_EMERALD

        p_a = tf.add_paragraph()
        p_a.text = p_auth
        p_a.font.size = Pt(11)
        p_a.font.color.rgb = TEXT_MUTED
        p_a.space_after = Pt(6)

        for bullet in p_bullets:
            pb = tf.add_paragraph()
            pb.text = f"•  {bullet}"
            pb.font.size = Pt(12)
            pb.font.color.rgb = TEXT_BODY
            pb.space_after = Pt(2)

    # -------------------------------------------------------------
    # SLIDE 6: Existing System
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide6)
    add_header(slide6, "Comparative Analysis", "Existing System & Critical Bottlenecks")

    exist_boxes = [
        ("Univariate Transaction Dependency", "Traditional systems (ARIMA, Holt-Winters) rely solely on historical order quantities. They are blind to emerging consumer sentiment and sudden shifts in market demand."),
        ("Neglect of Customer Review Text", "Standard ERP systems ignore millions of unstructured textual opinions, failing to capture qualitative customer perceptions, defect complaints, and praise."),
        ("Discrete Rating Bias & Fake Reviews", "Treats all 1–5 star ratings uniformly. False, sponsored, or low-quality reviews distort demand predictions because existing models lack helpfulness-based vote weighting."),
        ("High Forecast Error & Bullwhip Effect", "High forecast error (MAPE > 150%, extreme MSE) creates severe supply chain bullwhip distortions, causing over-stocking write-offs or catastrophic stockouts.")
    ]

    for idx, (title, desc) in enumerate(exist_boxes):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.7 + row * 2.6)

        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.35))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_ROSE
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p1 = tf.paragraphs[0]
        p1.text = f"❌  {title}"
        p1.font.bold = True
        p1.font.size = Pt(15)
        p1.font.color.rgb = ACCENT_ROSE
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_BODY

    # -------------------------------------------------------------
    # SLIDE 7: Proposed System
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide7)
    add_header(slide7, "Technical Innovation", "Proposed System: End-to-End AI Architecture")

    prop_boxes = [
        ("Dual-Engine Sentiment & Normalization", "Fuses BERT Transformer Softmax probabilities with VADER rule-based valence. Formulates Eq. 8 (0-1 mapping), Eq. 9 (1-5 sentiment scaling), and Eq. 10 (helpfulness normalization)."),
        ("SentiTSMixer Deep Time-Series Engine", "Replaces heavy recurrent/attention layers with Time-Mixing MLPs, Feature-Mixing MLPs, and 2D LayerNorm with residual connections. Features are injected at the Align Stage weighted by Random Forest MDI."),
        ("Aspect-Based Sentiment & Topic Discovery", "Deconstructs feedback into 6 key product aspects (Battery, Build, Price, UI, Speed, Support). Uncovers latent trends using BERTopic (UMAP + c-TF-IDF) and LDA topic distributions."),
        ("RAG AI Market Copilot & Decision Sandbox", "FAISS in-memory dense vector database enables instant similarity search over customer voices. Integrated with Gemini LLMs for strategic Q&A and What-If scenario simulation.")
    ]

    for idx, (title, desc) in enumerate(prop_boxes):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.7 + row * 2.6)

        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.35))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_EMERALD
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p1 = tf.paragraphs[0]
        p1.text = f"✅  {title}"
        p1.font.bold = True
        p1.font.size = Pt(15)
        p1.font.color.rgb = ACCENT_EMERALD
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_BODY

    # -------------------------------------------------------------
    # SLIDE 8: Problem Statement
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide8)
    add_header(slide8, "Research Focus", "Problem Statement")

    card8 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.1))
    card8.fill.solid()
    card8.fill.fore_color.rgb = CARD_BG
    card8.line.color.rgb = ACCENT_INDIGO
    card8.line.width = Pt(1)

    tf8 = card8.text_frame
    tf8.word_wrap = True
    tf8.margin_left = Inches(0.6)
    tf8.margin_right = Inches(0.6)
    tf8.margin_top = Inches(0.4)

    prob_statement_items = [
        "In modern multi-channel retail, consumer purchase decisions are heavily dictated by online ratings, detailed review narratives, and community upvotes.",
        "However, existing time-series sales forecasting models operate in isolation from unstructured textual data, creating a dangerous blind spot for demand planners.",
        "Furthermore, unweighted star ratings are vulnerable to fake reviews, extreme reviewer polarization, and discrete integer constraints that obscure subtle satisfaction shifts.",
        "Therefore, there is an urgent research and engineering challenge to design a unified AI framework that:",
        "1. Quantifies and normalizes multi-channel textual sentiment (BERT & VADER) on a continuous 1–5 scale.",
        "2. Dynamically balances community helpfulness votes to suppress unverified or false reviews.",
        "3. Amalgamates qualitative sentiment features into a high-speed, low-parameter deep forecasting model (SentiTSMixer) capable of slashing forecast error by 65% to 99%."
    ]

    for idx, item in enumerate(prob_statement_items):
        p = tf8.paragraphs[0] if idx == 0 else tf8.add_paragraph()
        p.text = item
        if idx == 3:
            p.font.bold = True
            p.font.size = Pt(15)
            p.font.color.rgb = ACCENT_CYAN
            p.space_before = Pt(8)
            p.space_after = Pt(6)
        elif idx > 3:
            p.font.size = Pt(14)
            p.font.color.rgb = TEXT_WHITE
            p.space_after = Pt(6)
        else:
            p.font.size = Pt(14)
            p.font.color.rgb = TEXT_BODY
            p.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 9: Objectives
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide9)
    add_header(slide9, "Project Goals", "Project Objectives")

    objectives = [
        ("Multi-Channel Data Pipeline", "Design and automate high-throughput ingestion from Twitter/X, Amazon Verified Reviews, and Reddit, resampled into weekly time-series matrices."),
        ("Dual-Engine Sentiment & Scale Normalization", "Implement BERT Transformer & VADER algorithms and formulate mathematical scaling (Eq. 8, 9, 10) to map sentiment and helpfulness to standard [1, 5] intervals."),
        ("Unsupervised Topic & Aspect Mining", "Extract latent consumer themes using Latent Dirichlet Allocation (LDA) and BERTopic (c-TF-IDF), alongside Aspect-Based Sentiment Analysis (ABSA) for 6 core product features."),
        ("SentiTSMixer Architecture Implementation", "Construct the SentiTSMixer deep forecasting network with Time-Mixing & Feature-Mixing MLPs, residual connections, and feature importance weighting at the Align Stage."),
        ("FAISS RAG & Conversational Copilot", "Build an in-memory dense vector index with FAISS for semantic similarity search over reviews, integrated with Gemini LLMs for market intelligence Q&A."),
        ("Interactive Multi-Dashboard Suite", "Deliver 8 enterprise-grade Streamlit dashboards for real-time monitoring, what-if simulations, supply chain bullwhip dampening, and automated report exports.")
    ]

    for idx, (title, desc) in enumerate(objectives):
        row = idx // 3
        col = idx % 3
        x = Inches(0.8 + col * 3.95)
        y = Inches(1.7 + row * 2.6)

        card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.8), Inches(2.35))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_CYAN
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)

        p1 = tf.paragraphs[0]
        p1.text = f"Goal {idx+1}: {title}"
        p1.font.bold = True
        p1.font.size = Pt(14)
        p1.font.color.rgb = ACCENT_CYAN
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_BODY

    # -------------------------------------------------------------
    # SLIDE 10: System Outputs & Visualizations (Part 1: Sentiment & ABSA)
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide10)
    add_header(slide10, "Experimental Results & Visualizations", "System Output (Part 1: Sentiment Pulse & Aspect Analysis)")

    img_p1 = "assets/outputs/output_1_sentiment_pulse.png"
    img_p2 = "assets/outputs/output_2_absa_radar.png"

    if os.path.exists(img_p1):
        slide10.shapes.add_picture(img_p1, Inches(0.8), Inches(1.65), width=Inches(5.75))
    if os.path.exists(img_p2):
        slide10.shapes.add_picture(img_p2, Inches(6.78), Inches(1.65), width=Inches(5.75))

    cap_box = slide10.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9))
    tf_cap = cap_box.text_frame
    tf_cap.word_wrap = True
    p_cap = tf_cap.paragraphs[0]
    p_cap.text = "Visual Findings: (Left) Real-time Net Sentiment Velocity Index across consumer sectors and multi-channel ingestion share. (Right) Granular Aspect-Based Sentiment Analysis (ABSA) polar radar chart isolating key hardware dimensions alongside 6-emotion distribution profiling."
    p_cap.font.size = Pt(11)
    p_cap.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 11: System Outputs & Visualizations (Part 2: Forecasting & Benchmarks)
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide11)
    add_header(slide11, "Experimental Results & Visualizations", "System Output (Part 2: Deep Demand Forecasting & Benchmarks)")

    img_p3 = "assets/outputs/output_3_demand_forecast.png"
    img_p4 = "assets/outputs/output_4_benchmarks.png"

    if os.path.exists(img_p3):
        slide11.shapes.add_picture(img_p3, Inches(0.8), Inches(1.65), width=Inches(5.75))
    if os.path.exists(img_p4):
        slide11.shapes.add_picture(img_p4, Inches(6.78), Inches(1.65), width=Inches(5.75))

    cap_box2 = slide11.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9))
    tf_cap2 = cap_box2.text_frame
    tf_cap2.word_wrap = True
    p_cap2 = tf_cap2.paragraphs[0]
    p_cap2.text = "Visual Findings: (Left) SentiTSMixer multi-step demand forecast with 95% confidence intervals tracking actual sales tightly across seasonal surges. (Right) Benchmark evaluation proving SentiTSMixer cuts MAPE to 12.99% (vs ARIMA's 248%), with the Cumulative Mixture contributing 48% MDI importance."
    p_cap2.font.size = Pt(11)
    p_cap2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 12: Conclusion
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide12)
    add_header(slide12, "Summary & Impact", "Conclusion & Future Work")

    c_col1 = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.75), Inches(5.1))
    c_col1.fill.solid()
    c_col1.fill.fore_color.rgb = CARD_BG
    c_col1.line.color.rgb = ACCENT_EMERALD
    c_col1.line.width = Pt(1)

    tf12_1 = c_col1.text_frame
    tf12_1.word_wrap = True
    tf12_1.margin_left = Inches(0.4)
    tf12_1.margin_right = Inches(0.4)
    tf12_1.margin_top = Inches(0.3)

    p = tf12_1.paragraphs[0]
    p.text = "🎯 Project Outcomes & Key Takeaways"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = ACCENT_EMERALD
    p.space_after = Pt(14)

    conclusions = [
        "Unprecedented Accuracy: SentiTSMixer achieves a 65% to 99% error reduction over ARIMA and LSTM models across varied dataset scales.",
        "Helpfulness Synergy: Factoring in community upvotes effectively neutralizes fake or misleading reviews, stabilizing demand variance.",
        "Feature Mixture Dominance: Random Forest MDI proves that the cumulative sentiment mixture (SC) provides up to 48%–91% of predictive influence.",
        "Enterprise Readiness: 8 interactive dashboards empower supply chain planners to anticipate demand shifts 3–6 weeks in advance."
    ]
    for c in conclusions:
        pb = tf12_1.add_paragraph()
        pb.text = f"•  {c}"
        pb.font.size = Pt(12)
        pb.font.color.rgb = TEXT_BODY
        pb.space_after = Pt(10)

    # Future Work
    c_col2 = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.7), Inches(5.75), Inches(5.1))
    c_col2.fill.solid()
    c_col2.fill.fore_color.rgb = CARD_BG
    c_col2.line.color.rgb = ACCENT_INDIGO
    c_col2.line.width = Pt(1)

    tf12_2 = c_col2.text_frame
    tf12_2.word_wrap = True
    tf12_2.margin_left = Inches(0.4)
    tf12_2.margin_right = Inches(0.4)
    tf12_2.margin_top = Inches(0.3)

    pf = tf12_2.paragraphs[0]
    pf.text = "🚀 Future Research & Extensions"
    pf.font.bold = True
    pf.font.size = Pt(16)
    pf.font.color.rgb = ACCENT_INDIGO
    pf.space_after = Pt(14)

    future_work = [
        "Automated Adversarial Fake Review Detection: Train zero-shot LLM classifiers to flag bot review farms before sentiment computation.",
        "Multilingual Sentiment Extraction: Expand Transformer tokenizers to process localized reviews across multilingual consumer markets.",
        "Real-Time WebSocket Streaming: Deploy Kafka / PubSub streaming pipelines for sub-second sentiment ingestion.",
        "Automated Procurement Triggers: Direct integration with enterprise ERP systems (SAP/Oracle) to autonomously adjust purchase orders based on What-If simulations."
    ]
    for fw in future_work:
        pb = tf12_2.add_paragraph()
        pb.text = f"•  {fw}"
        pb.font.size = Pt(12)
        pb.font.color.rgb = TEXT_BODY
        pb.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 13: References
    # -------------------------------------------------------------
    slide13 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide13)
    add_header(slide13, "Bibliography", "References")

    card13 = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.1))
    card13.fill.solid()
    card13.fill.fore_color.rgb = CARD_BG
    card13.line.color.rgb = ACCENT_CYAN
    card13.line.width = Pt(1)

    tf13 = card13.text_frame
    tf13.word_wrap = True
    tf13.margin_left = Inches(0.5)
    tf13.margin_right = Inches(0.5)
    tf13.margin_top = Inches(0.3)

    references = [
        "[1] P. Ghosh, S. Das, S. Roy, A. Bhattacharjee, A. Cortesi, and S. Sen, 'SentiTSMixer: A Specific Model for Sales Forecasting Using Sentiment Analysis of Customer,' IEEE Access, vol. 13, pp. 85882–85897, 2025.",
        "[2] S.-A. Chen, C.-L. Li, N. Yoder, S. O. Arik, and T. Pfister, 'TSMixer: An all-MLP architecture for time series forecasting,' arXiv:2303.06053, Google Research, 2023.",
        "[3] C. Hutto and É. Gilbert, 'VADER: A parsimonious rule-based model for sentiment analysis of social media text,' in Proc. 8th Int. AAAI Conf. Web Social Media (ICWSM), pp. 216–225, 2014.",
        "[4] J. Devlin, M.-W. Chang, K. Lee, and K. Toutanova, 'BERT: Pre-training of deep bidirectional transformers for language understanding,' in Proc. NAACL-HLT, pp. 4171–4186, 2019.",
        "[5] M. Grootendorst, 'BERTopic: Neural topic modeling with a class-based TF-IDF procedure,' arXiv:2203.05794, 2022.",
        "[6] J. Johnson, M. Douze, and H. Jégou, 'Billion-scale similarity search with GPUs [FAISS],' IEEE Transactions on Big Data, vol. 7, no. 3, pp. 535–547, 2021.",
        "[7] S. J. Taylor and B. Letham, 'Forecasting at scale [Prophet],' The American Statistician, vol. 72, no. 1, pp. 37–45, 2018.",
        "[8] P. Ghosh, O. Samanta, T. Goto, and S. Sen, 'Sales forecasting of overrated products: Fine tuning of customer's rating by integrating sentiment analysis,' IEEE Access, vol. 12, pp. 69578–69592, 2024."
    ]

    for idx, ref in enumerate(references):
        p = tf13.paragraphs[0] if idx == 0 else tf13.add_paragraph()
        p.text = ref
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_BODY
        p.space_after = Pt(8)

    prs.save(output_path)
    print(f"Presentation saved with embedded output images to {output_path}")

if __name__ == "__main__":
    create_presentation()
